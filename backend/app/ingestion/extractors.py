"""Text extraction for PDF, DOCX and TXT documents.

Scanned PDFs (no text layer) are detected and routed to an OCR provider when
one is configured; otherwise extraction fails loudly instead of silently
producing an empty document.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass

import filetype

from app.core.errors import InvalidInput, UnsupportedMedia

log = logging.getLogger(__name__)

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TXT = "text/plain"
SUPPORTED_MIME = {PDF, DOCX, TXT}
# Below this many characters per page, a PDF is treated as scanned (needs OCR)
MIN_CHARS_PER_PAGE = 40


@dataclass
class ExtractedText:
    text: str
    page_count: int | None
    method: str  # pdf-text | docx | txt | ocr:<provider>


class OcrProvider:
    """Interface for OCR services (Azure Document Intelligence, AWS Textract, Google Document AI)."""

    name = "none"

    def extract(self, data: bytes, mime_type: str) -> ExtractedText:  # pragma: no cover - interface
        raise NotImplementedError


def detect_mime(data: bytes, filename: str, declared: str | None) -> str:
    kind = filetype.guess(data)
    if kind is not None:
        if kind.mime == PDF:
            return PDF
        if kind.mime in (DOCX, "application/zip") and filename.lower().endswith(".docx"):
            return DOCX
        raise UnsupportedMedia(f"Unsupported file type: {kind.mime}")
    # No magic bytes: accept only if it decodes as text
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            data.decode("latin-1")
        except UnicodeDecodeError:  # pragma: no cover - latin-1 decodes everything
            raise UnsupportedMedia("Binary file of unknown type")
        if b"\x00" in data[:4096]:
            raise UnsupportedMedia("Binary file of unknown type")
    return TXT


def _pdf(data: bytes) -> ExtractedText:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise InvalidInput("PDF is password protected")
        pages = [(p.extract_text() or "") for p in reader.pages]
    except PdfReadError as exc:
        raise InvalidInput(f"Unreadable PDF: {exc}")
    text = "\n\n".join(pages)
    return ExtractedText(text=text, page_count=len(pages), method="pdf-text")


def _docx(data: bytes) -> ExtractedText:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ExtractedText(text="\n".join(parts), page_count=None, method="docx")


def _txt(data: bytes) -> ExtractedText:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return ExtractedText(text=text, page_count=None, method="txt")


def extract_text(data: bytes, mime_type: str, ocr: OcrProvider | None = None) -> ExtractedText:
    if mime_type == PDF:
        result = _pdf(data)
        pages = result.page_count or 1
        if len(result.text.strip()) < MIN_CHARS_PER_PAGE * pages:
            if ocr is None:
                raise InvalidInput(
                    "PDF has no usable text layer (scanned document). Configure an OCR provider "
                    "(Azure Document Intelligence / AWS Textract / Google Document AI)."
                )
            log.info("pdf_ocr_fallback", extra={"pages": pages, "provider": ocr.name})
            return ocr.extract(data, mime_type)
        return result
    if mime_type == DOCX:
        return _docx(data)
    if mime_type == TXT:
        return _txt(data)
    raise UnsupportedMedia(f"Unsupported MIME type {mime_type}")
